import io
import re
import zipfile
from collections import defaultdict

import pandas as pd
import pdfplumber
import streamlit as st


# =====================================================
# PAGE CONFIG
# =====================================================

st.set_page_config(
    page_title="Docket vs Billing Reconciler",
    page_icon="🧾",
    layout="wide",
)


# =====================================================
# DOCKET EXTRACTION
# (regex helpers copied unchanged from the A&N Docket Data Extractor)
# =====================================================

def normalise_material(raw_text):
    '''Maps raw material text to a canonical name (unchanged from extractor).'''
    if not raw_text:
        return ""

    lower = raw_text.lower()

    if "sand" in lower or "heid" in lower:
        return "Heidelberg Sand"

    if "roadbase" in lower or "road base" in lower or "road-base" in lower:
        return "Crushed Rock Basecourse"

    return ""


def extract_text_from_pdf(uploaded_file):
    '''Returns (page1_text, page2_text) for a docket PDF (unchanged).'''
    with pdfplumber.open(uploaded_file) as pdf:
        page1 = pdf.pages[0].extract_text() or "" if len(pdf.pages) > 0 else ""
        page2 = pdf.pages[1].extract_text() or "" if len(pdf.pages) > 1 else ""
    return page1, page2


def extract_docket_number(page1):
    '''Docket number, e.g. "A &-D-29868" (unchanged).'''
    match = re.search(r"Docket\s+(A\s*[&\-\w]+\s*D[\-\s]+\d+)", page1, re.IGNORECASE)
    if match:
        return match.group(1).strip()
    return ""


def extract_date(page1):
    '''First DD/MM/YYYY date on page 1 (unchanged).'''
    match = re.search(r"(\d{2}/\d{2}/\d{4})", page1)
    return match.group(1) if match else ""


def extract_material(page2):
    '''Canonical material name from page 2 (unchanged).'''
    lower = page2.lower()

    if "sand" in lower or "heild" in lower or "heid" in lower:
        return normalise_material("sand")

    if "roadbase" in lower or "road base" in lower or "road-base" in lower:
        return normalise_material("roadbase")

    return ""


def extract_total_tonnage(page2):
    '''Largest tonnage figure on page 2 = the docket total (unchanged).'''
    matches = re.findall(r"(\d+\.\d+)\s*T", page2)
    if not matches:
        return ""
    values = [float(x) for x in matches]
    return max(values)


def extract_individual_tonnages(page2):
    '''One float per "New Activity xx.xx T" load line (unchanged).'''
    matches = re.findall(r"New\s+Activity[\s\S]{0,10}?(\d+\.?\d*)\s*T", page2)
    return [float(x) for x in matches]


# ---------- NEW helpers (not in the original extractor) ----------

# The extractor calls the sand "Heidelberg Sand"; for reporting it is Gaskell Sand.
MATERIAL_LABELS = {"Heidelberg Sand": "Gaskell Sand"}


def material_label(canonical):
    return MATERIAL_LABELS.get(canonical, canonical)


def classify_billed_material(text):
    '''
    Maps the billing file's product text (e.g. "FILL SAND BELL RD") to
    "Gaskell Sand" or "Crushed Rock Basecourse". Returns "" if unrecognised.
    '''
    lower = str(text).lower()
    if "sand" in lower:
        return "Gaskell Sand"
    if any(k in lower for k in (
        "roadbase", "road base", "road-base", "basecourse", "base course",
        "crushed", "rock",
    )):
        return "Crushed Rock Basecourse"
    return ""


def normalise_plate(text):
    '''Upper-case and strip everything except letters/digits: "1jav 672" -> "1JAV672".'''
    return re.sub(r"[^A-Z0-9]", "", str(text).upper())


def extract_truck_rego(page1, known_plates=frozenset()):
    '''
    Reads the truck registration from the vehicle table on page 1:

        Truck Registration Trailer Registration Company Name Vehicle Type
        1jav672 1ucs508 H & S Bros pty ltd Tri-Axle

    Takes the first token on the line under the header. If that token is not one
    of the plates known from the billing file (e.g. the truck rego was left blank
    and the trailer rego came first), any token on that line that IS a known plate
    is preferred. As a last resort every token on page 1 is checked against the
    known plates. Returns "" if nothing is found.
    '''
    match = re.search(r"Truck\s+Registration[^\n]*\n\s*([^\n]+)", page1, re.IGNORECASE)
    if match:
        tokens = [normalise_plate(t) for t in match.group(1).split()]
        tokens = [t for t in tokens if t]
        for t in tokens:
            if t in known_plates:
                return t
        if tokens:
            return tokens[0]

    for t in page1.split():
        if normalise_plate(t) in known_plates:
            return normalise_plate(t)
    return ""


class NamedBytes(io.BytesIO):
    '''In-memory file with a .name, so PDFs from a ZIP behave like uploaded files.'''
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


def collect_pdfs(uploaded_files):
    '''Expands uploads into a flat list of PDF file-likes (ZIPs are opened).'''
    pdfs = []
    for f in uploaded_files:
        if f.name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(f.getvalue())) as zf:
                for info in zf.infolist():
                    base = info.filename.split("/")[-1]
                    if (
                        info.is_dir()
                        or not base.lower().endswith(".pdf")
                        or base.startswith(".")
                        or "__MACOSX" in info.filename
                    ):
                        continue
                    pdfs.append(NamedBytes(base, zf.read(info)))
        elif f.name.lower().endswith(".pdf"):
            pdfs.append(NamedBytes(f.name, f.getvalue()))
    return pdfs


def extract_docket_record(pdf_file, known_plates):
    '''Extracts everything needed for reconciliation from one docket PDF.'''
    try:
        page1, page2 = extract_text_from_pdf(pdf_file)
        date_str = extract_date(page1)
        return {
            "success": True,
            "pdf_name": pdf_file.name,
            "docket": extract_docket_number(page1),
            "date": pd.to_datetime(date_str, format="%d/%m/%Y", errors="coerce"),
            "date_str": date_str,
            "truck": extract_truck_rego(page1, known_plates),
            "material": material_label(extract_material(page2)),
            "total_tonnage": extract_total_tonnage(page2),
            "tonnages": extract_individual_tonnages(page2),
        }
    except Exception as e:
        return {"success": False, "pdf_name": pdf_file.name, "error": str(e)}


# =====================================================
# ORDER-DELIVERIES (BILLING) FILE
# =====================================================

# First non-blank of these gives the delivery's date/time (row can have blanks)
DATE_COLUMNS = [
    "Time leaving plant",
    "Time arrival on site",
    "Time start unload",
    "Time finish unload",
    "Done",
    "Batch started",
    "Time ticket printed",
    "Requested delivery date",
]
REQUIRED_COLUMNS = ["Docket", "License plate"]


def load_order_deliveries(file):
    '''
    Reads one order-deliveries export (.xlsx/.xls/.csv) into a tidy table with
    one row per load: Date, Time, Truck, Delivery Docket, Billed Tonnage, ...
    Cancelled / zero-quantity rows are split off into a separate "excluded" table.
    Returns (included_df, excluded_df).
    '''
    if file.name.lower().endswith(".csv"):
        raw = pd.read_csv(file)
    else:
        raw = pd.read_excel(file)

    raw.columns = [str(c).strip() for c in raw.columns]

    missing = [c for c in REQUIRED_COLUMNS if c not in raw.columns]
    qty_col = next((c for c in ("Load quantity", "Net weight") if c in raw.columns), None)
    if missing or qty_col is None:
        need = missing + ([] if qty_col else ["Load quantity"])
        raise ValueError(f"missing expected column(s): {', '.join(need)}")

    stamp = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
    for col in DATE_COLUMNS:
        if col in raw.columns:
            stamp = stamp.fillna(pd.to_datetime(raw[col], errors="coerce"))

    status = raw["Delivery status"].astype(str) if "Delivery status" in raw.columns else ""
    qty = pd.to_numeric(raw[qty_col], errors="coerce")

    text_cols = [c for c in ("Material", "Material description", "Customer material description")
                 if c in raw.columns]
    product_text = (
        raw[text_cols].astype(str).agg(" ".join, axis=1) if text_cols
        else pd.Series("", index=raw.index)
    )

    df = pd.DataFrame({
        "Date": stamp.dt.normalize(),
        "Time": stamp.dt.strftime("%H:%M"),
        "Truck": raw["License plate"].map(normalise_plate),
        "Material": product_text.map(classify_billed_material),
        "Delivery Docket": raw["Docket"].astype(str).str.replace(r"\.0$", "", regex=True),
        "Billed Tonnage": qty,
        "Delivery Status": status,
        "Product": raw[text_cols[0]] if text_cols else "",
        "Order": raw["Order number"] if "Order number" in raw.columns else "",
        "Source File": file.name,
    })

    is_cancelled = df["Delivery Status"].str.contains("cancel", case=False, na=False)
    is_empty = df["Billed Tonnage"].fillna(0) <= 0
    excluded = df[is_cancelled | is_empty].copy()
    included = df[~(is_cancelled | is_empty)].copy()
    return included, excluded


# =====================================================
# RECONCILIATION
# =====================================================

DIFF_COL = "What's different"
STATUS_MATCH = "Matched"
STATUS_DIFF = "Tonnage differs"
STATUS_REGO = "Rego differs"
STATUS_REGO_DIFF = "Rego & tonnage differ"
STATUS_BILLED_ONLY = "Billed - not on docket"
STATUS_DOCKET_ONLY = "On docket - not billed"
TRUCK_OK = "OK"
TRUCK_REGO = "Rego typo"
TRUCK_CHECK = "CHECK"

RED_STATUSES = {TRUCK_CHECK, STATUS_BILLED_ONLY, STATUS_DOCKET_ONLY}
AMBER_STATUSES = {STATUS_DIFF, STATUS_REGO, STATUS_REGO_DIFF, TRUCK_REGO}


def cents(x):
    return int(round(float(x) * 100))


def osa_distance(a, b):
    '''Edit distance where one swapped pair of neighbouring characters counts as 1.'''
    d = [[0] * (len(b) + 1) for _ in range(len(a) + 1)]
    for i in range(len(a) + 1):
        d[i][0] = i
    for j in range(len(b) + 1):
        d[0][j] = j
    for i in range(1, len(a) + 1):
        for j in range(1, len(b) + 1):
            cost = 0 if a[i - 1] == b[j - 1] else 1
            d[i][j] = min(d[i - 1][j] + 1, d[i][j - 1] + 1, d[i - 1][j - 1] + cost)
            if i > 1 and j > 1 and a[i - 1] == b[j - 2] and a[i - 2] == b[j - 1]:
                d[i][j] = min(d[i][j], d[i - 2][j - 2] + 1)
    return d[-1][-1]


def plates_similar(a, b):
    '''
    True if two different regos look like the same truck entered differently:
      - one is contained in the other  ("SW" in "SW25214"), or
      - they are one typo / one swapped pair apart ("1JAV672" vs "1JAV627").
    '''
    if not a or not b or a == b:
        return False
    short, long_ = sorted((a, b), key=len)
    if len(short) >= 2 and short in long_:
        return True
    return min(len(a), len(b)) >= 4 and osa_distance(a, b) <= 1


def reconcile(dockets, billed, pair_threshold, rego_threshold=2.0, match_similar_rego=True):
    '''
    Compares docket loads to billed loads. Pairing happens in this order:
      A. Same date + same rego, tonnage equal to the cent        -> "Matched"
      B. Same date, DIFFERENT but similar regos, tonnage within
         rego_threshold t (operator rego typo)                   -> "Rego differs" /
                                                                    "Rego & tonnage differ"
         Only attempted when at least one of the two regos does not appear on the
         other side that day (guards against sequential fleet numbers such as
         1IRP205 / 1IRP206, which are both real trucks).
      C. Same date + same rego, closest tonnage within
         pair_threshold t                                        -> "Tonnage differs"
      D. Whatever is left -> "Billed - not on docket" / "On docket - not billed"
    Returns (load_df, truck_df).
    '''
    d_loads = []
    for rec in dockets:
        for t in rec["tonnages"]:
            d_loads.append({
                "date": rec["date"], "truck": rec["truck"],
                "docket": rec["docket"] or rec["pdf_name"], "tonnage": t,
                "pdf": rec["pdf_name"], "material": rec["material"],
            })

    b_loads = []
    for _, r in billed.sort_values(["Date", "Time"], na_position="last").iterrows():
        b_loads.append({
            "date": r["Date"], "truck": r["Truck"], "delivery": r["Delivery Docket"],
            "time": r["Time"], "tonnage": float(r["Billed Tonnage"]),
            "source": r["Source File"], "material": r["Material"],
        })

    d_by_key, b_by_key = defaultdict(list), defaultdict(list)
    for i, l in enumerate(d_loads):
        d_by_key[(l["date"], l["truck"])].append(i)
    for i, l in enumerate(b_loads):
        b_by_key[(l["date"], l["truck"])].append(i)

    pairs = []  # (docket idx, billed idx, kind)
    free_d, free_b = set(range(len(d_loads))), set(range(len(b_loads)))

    def take(di, bi, kind):
        pairs.append((di, bi, kind))
        free_d.discard(di)
        free_b.discard(bi)

    # A. exact tonnage, same rego
    for key, dis in d_by_key.items():
        for di in dis:
            for bi in b_by_key.get(key, []):
                if bi in free_b and cents(d_loads[di]["tonnage"]) == cents(b_loads[bi]["tonnage"]):
                    take(di, bi, "exact")
                    break

    # B. similar (but different) regos, tonnage within rego_threshold
    if match_similar_rego:
        d_plates, b_plates = defaultdict(set), defaultdict(set)
        for l in d_loads:
            d_plates[l["date"]].add(l["truck"])
        for l in b_loads:
            b_plates[l["date"]].add(l["truck"])

        cands = []
        for di in free_d:
            dl = d_loads[di]
            for bi in free_b:
                bl = b_loads[bi]
                if dl["date"] != bl["date"] or dl["truck"] == bl["truck"]:
                    continue
                diff = abs(dl["tonnage"] - bl["tonnage"])
                if diff > rego_threshold + 1e-9 or not plates_similar(dl["truck"], bl["truck"]):
                    continue
                orphan = (
                    dl["truck"] not in b_plates[dl["date"]]
                    or bl["truck"] not in d_plates[dl["date"]]
                )
                if orphan:
                    cands.append((round(diff, 4), di, bi))
        for diff, di, bi in sorted(cands):
            if di in free_d and bi in free_b:
                take(di, bi, "rego")

    # C. same rego, closest tonnage within pair_threshold
    for key, dis in d_by_key.items():
        cands = sorted(
            (abs(d_loads[di]["tonnage"] - b_loads[bi]["tonnage"]), di, bi)
            for di in dis if di in free_d
            for bi in b_by_key.get(key, []) if bi in free_b
        )
        for diff, di, bi in cands:
            if diff <= pair_threshold and di in free_d and bi in free_b:
                take(di, bi, "near")

    # ---- build load rows ----
    def make_row(di, bi, status, note):
        dl = d_loads[di] if di is not None else None
        bl = b_loads[bi] if bi is not None else None
        docket_t = dl["tonnage"] if dl else None
        billed_t = bl["tonnage"] if bl else None
        if dl and bl:
            diff = round(billed_t - docket_t, 2)
        else:
            diff = billed_t if bl else -docket_t
        return {
            "Date": (bl or dl)["date"],
            "Truck": bl["truck"] if bl else dl["truck"],
            "Material": (dl["material"] if dl and dl["material"] else (bl["material"] if bl else "")),
            "Docket Rego": dl["truck"] if dl else "",
            "Docket": dl["docket"] if dl else "",
            "Docket Tonnage": docket_t,
            "Billed Delivery Docket": bl["delivery"] if bl else "",
            "Billed Time": bl["time"] if bl else "",
            "Billed Tonnage": billed_t,
            "Billed - Docket (t)": diff,
            "Status": status,
            "Note": note,
            "Source File": bl["source"] if bl else "",
            "Docket PDF": dl["pdf"] if dl else "",
        }

    def sort_key(row, group, idx):
        date = pd.Timestamp.max if pd.isna(row["Date"]) else row["Date"]
        return (date, row["Truck"], group, idx)

    keyed = []
    for di, bi, kind in pairs:
        note = ""
        if kind == "exact":
            status = STATUS_MATCH
        elif kind == "near":
            status = STATUS_DIFF
        else:
            same = cents(d_loads[di]["tonnage"]) == cents(b_loads[bi]["tonnage"])
            status = STATUS_REGO if same else STATUS_REGO_DIFF
            note = (f"Docket rego '{d_loads[di]['truck']}' is similar to billed rego "
                    f"'{b_loads[bi]['truck']}' - likely entered incorrectly")
        row = make_row(di, bi, status, note)
        keyed.append((sort_key(row, 0, di), row))
    for bi in free_b:
        bl = b_loads[bi]
        note = "" if d_by_key.get((bl["date"], bl["truck"])) else \
            "No docket uploaded for this truck on this date"
        row = make_row(None, bi, STATUS_BILLED_ONLY, note)
        keyed.append((sort_key(row, 1, bi), row))
    for di in free_d:
        dl = d_loads[di]
        note = "" if b_by_key.get((dl["date"], dl["truck"])) else \
            "Truck not in billing file for this date"
        row = make_row(di, None, STATUS_DOCKET_ONLY, note)
        keyed.append((sort_key(row, 2, di), row))

    load_df = pd.DataFrame([r for _, r in sorted(keyed, key=lambda x: x[0])])
    if load_df.empty:
        return load_df, pd.DataFrame()

    # ---- one row per truck per day ----
    stated_notes = {}
    for rec in dockets:
        loads_sum = round(sum(rec["tonnages"]), 2)
        stated = rec["total_tonnage"] if isinstance(rec["total_tonnage"], (int, float)) else loads_sum
        if abs(stated - loads_sum) > 0.005:
            stated_notes[rec["docket"] or rec["pdf_name"]] = (
                f"{rec['docket']}: stated total {stated:.2f} t but loads add to {loads_sum:.2f} t"
            )

    truck_rows = []
    for (date, truck), g in load_df.groupby(["Date", "Truck"], sort=False, dropna=False):
        docket_sum = round(g["Docket Tonnage"].sum(), 2)
        billed_sum = round(g["Billed Tonnage"].sum(), 2)
        dockets_here = [x for x in dict.fromkeys(g["Docket"]) if x]
        notes = [stated_notes[x] for x in dockets_here if x in stated_notes]
        bad = g[g["Status"] != STATUS_MATCH]
        kinds = set(bad["Status"])
        if (kinds - {STATUS_REGO}) or notes:
            status = TRUCK_CHECK
        elif kinds:
            status = TRUCK_REGO
        else:
            status = TRUCK_OK
        truck_rows.append({
            "Date": date,
            "Truck": truck,
            "Material": ", ".join(dict.fromkeys(m for m in g["Material"] if m)),
            "Dockets": ", ".join(dockets_here) or "(none)",
            "Docket Loads": int(g["Docket Tonnage"].notna().sum()),
            "Docket Total (t)": docket_sum,
            "Billed Loads": int(g["Billed Tonnage"].notna().sum()),
            "Billed Total (t)": billed_sum,
            "Billed - Docket (t)": round(billed_sum - docket_sum, 2),
            "Status": status,
            DIFF_COL: "; ".join(
                [f"{s}: {n}" for s, n in bad["Status"].value_counts().items()] + notes
            ),
        })
    return load_df, pd.DataFrame(truck_rows)


# =====================================================
# EXCEL REPORT
# =====================================================

def build_report(truck_df, load_df, docket_df, excluded_df):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    issues_df = load_df[load_df["Status"] != STATUS_MATCH]

    sheets = {
        "Summary by Truck": truck_df,
        "Load Comparison": load_df,
        "Issues Only": issues_df,
        "Docket Extraction": docket_df,
        "Excluded Billing Rows": excluded_df,
    }

    with pd.ExcelWriter(buffer, engine="openpyxl", date_format="DD/MM/YYYY") as writer:
        for name, df in sheets.items():
            out = df.copy()
            if "Date" in out.columns:
                out["Date"] = pd.to_datetime(out["Date"]).dt.date
            out.to_excel(writer, sheet_name=name, index=False)

        red = PatternFill("solid", fgColor="F8CBAD")
        amber = PatternFill("solid", fgColor="FFE699")
        for name, ws in writer.sheets.items():
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col_cells in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(width + 3, 60)

            headers = [c.value for c in ws[1]]
            if "Status" in headers:
                idx = headers.index("Status") + 1
                for row in ws.iter_rows(min_row=2):
                    val = row[idx - 1].value
                    fill = None
                    if val in RED_STATUSES:
                        fill = red
                    elif val in AMBER_STATUSES:
                        fill = amber
                    if fill:
                        for c in row:
                            c.fill = fill

    buffer.seek(0)
    return buffer.getvalue()


# =====================================================
# UI HELPERS
# =====================================================

def style_status(df):
    def colour(row):
        s = row.get("Status")
        if s in RED_STATUSES:
            return ["background-color: #f8cbad; color: #000"] * len(row)
        if s in AMBER_STATUSES:
            return ["background-color: #ffe699; color: #000"] * len(row)
        return [""] * len(row)

    num_cols = [c for c in df.columns if "(t)" in c or "Tonnage" in c]
    styler = df.style.apply(colour, axis=1).format("{:.2f}", subset=num_cols, na_rep="")
    if "Date" in df.columns:
        styler = styler.format({"Date": lambda d: "" if pd.isna(d) else d.strftime("%d/%m/%Y")})
    return styler


def render_results(res):
    truck_df, load_df = res["truck_df"], res["load_df"]

    docket_total = truck_df["Docket Total (t)"].sum()
    billed_total = truck_df["Billed Total (t)"].sum()
    n_issues = int((load_df["Status"] != STATUS_MATCH).sum())

    c = st.columns(5)
    c[0].metric("Docket total (t)", f"{docket_total:,.2f}")
    c[1].metric("Billed total (t)", f"{billed_total:,.2f}")
    c[2].metric("Billed − Docket (t)", f"{billed_total - docket_total:+,.2f}")
    c[3].metric("Loads matched", int((load_df["Status"] == STATUS_MATCH).sum()))
    c[4].metric("Loads needing review", n_issues)

    n_rego = int(load_df["Status"].isin([STATUS_REGO, STATUS_REGO_DIFF]).sum())
    if n_rego:
        st.info(
            f"{n_rego} load(s) were paired despite a rego mismatch between docket and billing "
            "(likely operator entry errors). They are not counted as missing loads but are "
            "listed below so the rego can be corrected."
        )

    if n_issues == 0 and (truck_df["Status"] == TRUCK_OK).all():
        st.success("Everything reconciles - no discrepancies found.")
    else:
        st.error(f"{n_issues} load(s) need review. Details below.")

    st.subheader("Issues")
    issues = load_df[load_df["Status"] != STATUS_MATCH]
    if issues.empty:
        st.write("No load-level issues.")
    else:
        st.dataframe(style_status(issues), use_container_width=True, hide_index=True)

    notes = truck_df[truck_df[DIFF_COL].str.contains("stated total", na=False)]
    for _, r in notes.iterrows():
        st.warning(f"{r['Truck']}: {r[DIFF_COL]}")

    st.subheader("Summary by truck")
    st.dataframe(style_status(truck_df), use_container_width=True, hide_index=True)

    with st.expander("All loads, side by side", expanded=False):
        st.dataframe(style_status(load_df), use_container_width=True, hide_index=True)

    with st.expander("Docket extraction (what was read from each PDF)", expanded=False):
        st.dataframe(res["docket_df"], use_container_width=True, hide_index=True)

    if not res["excluded_df"].empty:
        with st.expander(f"Excluded billing rows - cancelled / zero quantity ({len(res['excluded_df'])})"):
            st.dataframe(res["excluded_df"], use_container_width=True, hide_index=True)

    st.download_button(
        "Download reconciliation report (Excel)",
        data=res["report"],
        file_name="Docket_Reconciliation.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# =====================================================
# MAIN APP
# =====================================================

def main():
    st.title("Docket vs Billing Reconciler")
    st.markdown(
        "Compare the tonnages on your **A&N dockets** against the **order-deliveries** "
        "billing export. Loads are matched by date + truck rego, then by tonnage."
    )

    with st.sidebar:
        st.header("Settings")
        pair_threshold = st.number_input(
            "Max tonnage gap to treat as the same load (t)",
            min_value=0.0, max_value=50.0, value=5.0, step=0.5,
            help="A leftover docket load and billed load on the same truck/day are reported "
                 "as 'Tonnage differs' if they are within this many tonnes of each other. "
                 "Otherwise they are listed as separate missing loads.",
        )
        match_similar_rego = st.checkbox(
            "Match loads with similar regos (operator typos)", value=True,
            help="e.g. 'SW' on the docket vs 'SW25214' in billing.",
        )
        rego_threshold = st.number_input(
            "Max tonnage gap for similar-rego matches (t)",
            min_value=0.0, max_value=50.0, value=2.0, step=0.5,
            disabled=not match_similar_rego,
        )

    left, right = st.columns(2)
    with left:
        billed_files = st.file_uploader(
            "1. Order-deliveries files (billing export)",
            type=["xlsx", "xls", "csv"],
            accept_multiple_files=True,
        )
    with right:
        docket_files = st.file_uploader(
            "2. A&N docket PDFs (or a ZIP of PDFs)",
            type=["pdf", "zip"],
            accept_multiple_files=True,
        )

    if st.button("Compare", type="primary", disabled=not (billed_files and docket_files)):
        problems = []

        # ---- billing files ----
        included, excluded = [], []
        for f in billed_files:
            try:
                inc, exc = load_order_deliveries(f)
                included.append(inc)
                excluded.append(exc)
            except Exception as e:
                problems.append(f"Could not read **{f.name}**: {e}")
        if problems:
            for p in problems:
                st.error(p)
            st.stop()

        billed = pd.concat(included, ignore_index=True)
        excluded_df = pd.concat(excluded, ignore_index=True)
        known_plates = frozenset(billed["Truck"].unique())

        unknown_products = sorted(set(billed.loc[billed["Material"] == "", "Product"].astype(str)))
        if unknown_products:
            st.warning(
                "Couldn't classify the material for billed product(s): "
                + ", ".join(unknown_products) + ". Their Material cell will be blank."
            )

        # ---- docket PDFs ----
        pdfs = collect_pdfs(docket_files)
        if not pdfs:
            st.error("No PDF files found in the docket uploads.")
            st.stop()

        records, seen = [], set()
        progress = st.progress(0)
        for i, pdf in enumerate(pdfs):
            rec = extract_docket_record(pdf, known_plates)
            if not rec["success"]:
                st.warning(f"Could not read **{rec['pdf_name']}**: {rec['error']}")
            elif rec["docket"] and rec["docket"] in seen:
                st.warning(f"Duplicate docket {rec['docket']} ({rec['pdf_name']}) ignored.")
            else:
                seen.add(rec["docket"])
                records.append(rec)
                if not rec["truck"]:
                    st.warning(f"No truck rego found in **{rec['pdf_name']}** - its loads can't be matched.")
                if pd.isna(rec["date"]):
                    st.warning(f"No date found in **{rec['pdf_name']}**.")
                if not rec["material"]:
                    st.warning(f"Material not detected in **{rec['pdf_name']}**.")
                if not rec["tonnages"]:
                    st.warning(f"No load tonnages found in **{rec['pdf_name']}**.")
            progress.progress((i + 1) / len(pdfs))
        progress.empty()

        if not records:
            st.error("No dockets could be read.")
            st.stop()

        load_df, truck_df = reconcile(
            records, billed, pair_threshold, rego_threshold, match_similar_rego
        )

        docket_df = pd.DataFrame([{
            "PDF": r["pdf_name"], "Docket": r["docket"], "Date": r["date"],
            "Truck": r["truck"], "Material": r["material"],
            "Loads": len(r["tonnages"]), "Stated Total (t)": r["total_tonnage"],
            "Sum of Loads (t)": round(sum(r["tonnages"]), 2),
        } for r in records])

        st.session_state["result"] = {
            "truck_df": truck_df,
            "load_df": load_df,
            "docket_df": docket_df,
            "excluded_df": excluded_df,
            "report": build_report(truck_df, load_df, docket_df, excluded_df),
        }

    # Results live in session_state so they survive the rerun caused by downloading
    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
